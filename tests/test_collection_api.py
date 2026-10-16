"""API → real background thread → durable store → CSV/XLSX and recovery."""
import asyncio
import base64
import csv
import io
from datetime import datetime
from unittest.mock import AsyncMock
import httpx
import pytest
from openpyxl import load_workbook
from src.api import main as api
from src.collection import engine
from src.collection.api import _active
from src.collection.store import Store
from src.crawlers.twitter_url import TwitterCookieFetcher


@pytest.fixture(autouse=True)
def isolate(tmp_path,monkeypatch):
    monkeypatch.setenv('VOCAB_DATA_DIR',str(tmp_path/'data'))
    engine.Traffic.active.clear();engine.Traffic.delay.clear();engine.Traffic.until.clear()


async def done(client,rid):
    for _ in range(500):
        state=(await client.get('/api/task-status',params={'task_id':rid})).json()
        if state['status']!='running' and rid not in _active:
            return state
        await asyncio.sleep(.01)
    raise AssertionError('background task did not exit')


async def test_api_resume_only_comments_and_export_after_memory_cache_cleared(init_test_db,monkeypatch):
    search=AsyncMock(return_value=([{'tweet_id':'p','content':'private text','url':'https://x.com/a/status/p','media_urls':'https://example/p.jpg'}],''))
    monkeypatch.setattr(TwitterCookieFetcher,'search_tweets',search)
    comment_calls=[]
    async def comments(fetcher,platform,posts,cookies):
        comment_calls.append(cookies['auth_token'])
        posts[0].update(comments_status='finished',replies_data=[{'tweet_id':'c','content':'reply'}])
        return posts
    monkeypatch.setattr(engine,'fetch_comments',comments)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),base_url='http://local') as client:
        response=await client.post('/api/search',json={'keyword':'kw','count':1,'cookies':[{'platform':'twitter','ct0':'c','auth_token':'old'}]})
        rid=response.json()['task_id']; first=await done(client,rid)
        assert first['result']['total_posts']==1
        await api._save_platform_cookies_db('twitter',{'ct0':'new','auth_token':'new'})
        assert (await client.post('/api/collections/'+rid+'/resume?comments_only=true')).status_code==200
        resumed=await done(client,rid)
        assert search.await_count==1 and comment_calls==['new']
        assert resumed['result']['total_comments']==1
        with api._BG_TASKS_LOCK:api._BACKGROUND_TASKS.clear()
        summary=(await client.get('/api/collections')).json()
        assert 'private text' not in str(summary)
        exported=(await client.get('/api/collections/'+rid+'/result')).json()
        rows=list(csv.DictReader(io.StringIO(base64.b64decode(exported['csv_data']).decode('utf-8-sig'))))
        assert rows[1]['parent_id']=='p'
        wb=load_workbook(io.BytesIO(base64.b64decode(exported['xlsx_data'])))
        assert wb.active.max_row==3;wb.close()
        assert exported['sampled_posts']==[]


async def test_preflight_empty_search_is_accessible_not_auth_failure(init_test_db,monkeypatch):
    monkeypatch.setattr(TwitterCookieFetcher,'search_tweets',AsyncMock(return_value=([],'')))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),base_url='http://local') as client:
        response=await client.post('/api/collection-preflight',json={'platforms':['twitter'],'cookies':[{'platform':'twitter','ct0':'c','auth_token':'a'}]})
        state=await done(client,response.json()['task_id'])
        assert state['result']['preflight']
        assert state['result']['status']=='success'
        assert not state['result']['errors']


async def test_invalid_schedule_never_saved_and_valid_schedule_preview_matches(init_test_db,monkeypatch):
    monkeypatch.setattr(api,'_reload_scheduler',lambda:None)
    from src.common.database import get_db
    db=await get_db()
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),base_url='http://local') as client:
        body={'name':'weekly','cron_expression':'0 15 * * 8','params':{'keywords':['test'],'platforms':['twitter']}}
        invalid=(await client.post('/api/scheduled-tasks',json=body)).json()
        assert invalid['status']=='error'
        assert (await (await db.execute('SELECT COUNT(*) FROM scheduled_tasks')).fetchone())[0]==0
        body['cron_expression']='0 15 * * 3'
        good=(await client.post('/api/scheduled-tasks',json=body)).json()
        assert good['status']=='ok'
        task=(await client.get('/api/scheduled-tasks')).json()['tasks'][0]
        assert task['description']=='周三 15:00'
        assert not task['legacy_weekdays'] and task['timezone']=='Asia/Shanghai'


async def test_scheduled_url_uses_same_durable_engine(init_test_db,monkeypatch,tmp_path):
    from src.common.database import get_db
    db=await get_db();stamp=datetime.now().isoformat()
    await db.execute('INSERT INTO scheduled_tasks(id,created_at,updated_at) VALUES(?,?,?)',('profile',stamp,stamp));await db.commit()
    await api._save_platform_cookies_db('twitter',{'ct0':'c','auth_token':'a'})
    fetch=AsyncMock(return_value=([{'tweet_id':'post','content':'profile post'}],''))
    monkeypatch.setattr(TwitterCookieFetcher,'fetch_user_tweets',fetch)
    await api._execute_scheduled_task({'id':'profile','name':'主页任务','task_type':'url_fetch','save_path':str(tmp_path),
        'params':{'urls':['https://x.com/example'],'platforms':['twitter'],'count':1,'analyze':False}})
    assert fetch.call_args.args[0]==['example']
    run=Store().list(owner='profile')[0]
    assert run['status']=='success'
    assert Store().rows(run['id'])[0]['content']=='profile post'
    saved_files = list(tmp_path.glob('*.csv'))
    assert len(saved_files) == 1 and saved_files[0].stem.startswith('主页任务')
    # 重新打开数据库、修改运行更新时间后，下载仍使用原来固定的完成日期。
    store = Store()
    with store.connect() as connection:
        connection.execute("UPDATE runs SET updated='2030-01-01T00:00:00+00:00' WHERE id=?", (run['id'],))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=api.app),base_url='http://local') as client:
        exported = (await client.get('/api/collections/'+run['id']+'/result')).json()
    assert exported['csv_filename'] == saved_files[0].name
    assert exported['xlsx_filename'] == saved_files[0].with_suffix('.xlsx').name


async def test_legacy_scheduled_analysis_is_ignored(init_test_db,monkeypatch,tmp_path):
    from src.common.database import get_db
    from src.orchestrator.pipeline import Pipeline
    db=await get_db();stamp=datetime.now().isoformat()
    await db.execute('INSERT INTO scheduled_tasks(id,created_at,updated_at) VALUES(?,?,?)',('audit',stamp,stamp));await db.commit()
    await api._save_platform_cookies_db('twitter',{'ct0':'c','auth_token':'a'})
    monkeypatch.setattr(TwitterCookieFetcher,'search_tweets',AsyncMock(return_value=([{'tweet_id':'post','content':'saved real data'}],'')))
    mock_audit=AsyncMock(side_effect=AssertionError('mock results must not enter candidate database'))
    monkeypatch.setattr(Pipeline,'process_posts',mock_audit)
    await db.execute("INSERT INTO model_config(config_key,config_value,updated_at) VALUES('model_api_key','legacy-secret','2026-10-15')")
    await db.execute("INSERT INTO model_config(config_key,config_value,updated_at) VALUES('model_base_url','https://example.test/retired-model','2026-10-15')")
    await db.commit()
    await api._execute_scheduled_task({'id':'audit','save_path':str(tmp_path),
        'params':{'keywords':['a'],'platforms':['twitter'],'count':1,'analyze':True}})
    mock_audit.assert_not_called()
    state=await (await db.execute('SELECT last_run_status,last_error FROM scheduled_tasks WHERE id=?',('audit',))).fetchone()
    assert state['last_run_status']=='success'
    assert not state['last_error']
    assert (await (await db.execute('SELECT COUNT(*) FROM vocabulary')).fetchone())[0]==0
    assert list(tmp_path.glob('*.csv'))
